from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Local_Multiplayer_Feasibility_Report.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(30, 35, 40)
MUTED = RGBColor(95, 103, 112)
LIGHT = "F2F4F7"
PALE_BLUE = "EAF2F8"
PALE_GREEN = "EAF5EE"
PALE_AMBER = "FFF5D9"
WHITE = "FFFFFF"
USABLE_DXA = 9360


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D7DBE2", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:cantSplit")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def create_decimal_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids or [0]) + 1
    num_id = max(num_ids or [0]) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    nsid = OxmlElement("w:nsid")
    nsid.set(qn("w:val"), f"A1{abstract_id:06X}"[-8:])
    tmpl = OxmlElement("w:tmpl")
    tmpl.set(qn("w:val"), f"B2{abstract_id:06X}"[-8:])
    abstract.extend([nsid, tmpl])
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "decimal")
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "%1.")
    justify = OxmlElement("w:lvlJc")
    justify.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    p_pr.extend([tabs, indent])
    level.extend([start, fmt, text, justify, p_pr])
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def create_bullet_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids or [0]) + 1
    num_id = max(num_ids or [0]) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    nsid = OxmlElement("w:nsid")
    nsid.set(qn("w:val"), f"C3{abstract_id:06X}"[-8:])
    tmpl = OxmlElement("w:tmpl")
    tmpl.set(qn("w:val"), f"D4{abstract_id:06X}"[-8:])
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.extend([nsid, tmpl, multi])
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "bullet")
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "\u2022")
    justify = OxmlElement("w:lvlJc")
    justify.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    p_pr.append(indent)
    level.extend([start, fmt, text, justify, p_pr])
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_ref = OxmlElement("w:numId")
    num_ref.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_ref])
    p_pr.insert(0, num_pr)


def set_run(run, size=11, bold=False, italic=False, color=INK, font="Calibri"):
    run.font.name = font
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    return run


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2E74B5")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, result, end])
    set_run(run, size=9, color=MUTED)


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:color"), "2E74B5")
    left.set(qn("w:space"), "6")
    borders.append(left)
    p_pr.append(borders)


def add_body(doc, text, bold_lead=None, italic=False, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if bold_lead and text.startswith(bold_lead):
        set_run(p.add_run(bold_lead), bold=True)
        set_run(p.add_run(text[len(bold_lead):]), italic=italic)
    else:
        set_run(p.add_run(text), italic=italic)
    return p


def add_bullet(doc, text, level=0, bold_lead=None):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Inches(0.5 if level == 0 else 0.75)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.167
    p.paragraph_format.keep_together = True
    if bold_lead and text.startswith(bold_lead):
        set_run(p.add_run(bold_lead), bold=True)
        set_run(p.add_run(text[len(bold_lead):]))
    else:
        set_run(p.add_run(text))
    return p


def add_number(doc, text, num_id):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    p.paragraph_format.keep_together = True
    set_run(p.add_run(text))
    return p


def add_heading(doc, text, level=1, page_break_before=False):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break_before
    return p


def add_callout(doc, label, text, fill=PALE_BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.10
    shade_paragraph(p, fill)
    set_run(p.add_run(label + "  "), bold=True, color=DARK_BLUE)
    set_run(p.add_run(text))
    return p


def add_table(doc, headers, rows, widths_dxa, font_size=9.4):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header = table.rows[0]
    repeat_header(header)
    for i, title in enumerate(headers):
        set_cell_fill(header.cells[i], LIGHT)
        p = header.cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(title), size=font_size, bold=True, color=DARK_BLUE)
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            p = row.cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            set_run(p.add_run(str(value)), size=font_size)
    set_table_geometry(table, widths_dxa)
    for row in table.rows:
        prevent_row_split(row)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
    specs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_document(doc):
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    doc.settings.odd_and_even_pages_header_footer = False
    header = section.header
    p = header.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run("THE PICTURE SHOP  |  TECHNICAL FEASIBILITY"), size=8.5, bold=True, color=MUTED)
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run("Local Multiplayer Feasibility Report  |  "), size=9, color=MUTED)
    add_field(p, "PAGE")


def build():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(4)
    set_run(title.add_run("LOCAL MULTIPLAYER FEASIBILITY REPORT"), size=24, bold=True, color=INK)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    set_run(subtitle.add_run("The Picture Shop mobile version - same-area cooperative play over a local network"), size=13.5, color=MUTED)
    for label, value in (
        ("Prepared for", "The Picture Shop development team"),
        ("Assessment date", "August 25, 2026"),
        ("Current platform", "Android, LÖVE 11.5, shared Lua game source"),
        ("Recommended launch scope", "2-4 players; same Wi-Fi or phone hotspot; host-authoritative shared shop"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        set_run(p.add_run(label + ": "), bold=True)
        set_run(p.add_run(value))

    add_callout(
        doc,
        "Decision",
        "Local multiplayer is feasible. The strongest first release is Wi-Fi LAN/hotspot play using LÖVE's included ENet transport, with one phone acting as the authoritative host. Bluetooth and Wi-Fi Direct should not be the first implementation.",
        PALE_GREEN,
    )

    add_heading(doc, "1. Executive summary")
    add_body(doc, "The requested experience is achievable without an internet server: players in the same room can connect through the same Wi-Fi router, or one device can provide a hotspot and the others can join it. The game should treat one device as the host and source of truth for the shop, while guest devices send player intentions and receive replicated state.")
    add_body(doc, "This approach fits the project's existing architecture better than Bluetooth. The current Android build already contains LÖVE's ENet and LuaSocket networking modules, but the app manifest explicitly removes Android's INTERNET permission. The main work is therefore not adding a radio API; it is converting the current single-player global state into a safe host-authoritative multiplayer model and adding lobby, discovery, synchronization, conflict rules, and network tests.")
    add_bullet(doc, "Recommended transport: ENet over UDP for the game session; LuaSocket UDP broadcast for initial LAN discovery.", bold_lead="Recommended transport:")
    add_bullet(doc, "Recommended connection modes: shared Wi-Fi first, user-created phone hotspot as the no-router fallback.", bold_lead="Recommended connection modes:")
    add_bullet(doc, "Recommended session size: 2-4 players for the first release.", bold_lead="Recommended session size:")
    add_bullet(doc, "Recommended authority model: host simulates time, money, inventory, machines, jobs, visitors, trucks, and saves.", bold_lead="Recommended authority model:")
    add_bullet(doc, "Deferred: host migration, internet matchmaking, Bluetooth transport, Wi-Fi Direct, cross-platform iOS certification, and adversarial anti-cheat.", bold_lead="Deferred:")

    add_heading(doc, "2. What was found in the current project")
    add_table(doc, ["Finding", "Evidence in this repository", "Multiplayer consequence"], [
        ("One shared codebase", "main.lua, conf.lua, and src/ are shared by Windows and Android.", "Networking should remain shared Lua code; Android-specific work should stay at the wrapper edge."),
        ("Single authoritative state today", "src/app.lua owns one State.new() instance and advances calendar, world, machines, email, and visitors locally.", "Guests cannot run the same simulation independently without divergence."),
        ("Single player singleton", "src/world.lua stores one World.player and saves only one x/y position.", "Introduce participants[playerId] and render multiple avatars."),
        ("Local-only persistence", "src/save.lua writes versioned Lua tables; save schema is version 12.", "Only the host should write the shared shop save. Network payloads need a separate safe codec."),
        ("Frequent direct mutation", "Input and screen modules call domain functions that mutate state and often save immediately.", "Route multiplayer actions through validated commands handled by the host."),
        ("Networking libraries present", "The staged LÖVE 11.5 Android source preloads both enet and socket.", "No new native transport library is required for LAN play."),
        ("Networking permission removed", "mobile/android/AndroidManifest.xml removes android.permission.INTERNET.", "LAN sockets will fail until this permission is restored."),
    ], [2100, 3480, 3780], font_size=8.8)

    add_callout(doc, "Important security boundary", "Do not reuse src/save.lua's loadstring-based local save parser for network messages. Network data must use a non-executable format with strict type, size, range, and version validation.", PALE_AMBER)

    add_heading(doc, "3. Connection technology comparison")
    add_table(doc, ["Option", "Fit", "Advantages", "Costs / limitations", "Recommendation"], [
        ("Same Wi-Fi LAN", "Excellent", "Fast, low latency, works with ENet/LuaSocket, supports 2-4+ peers, shared Lua implementation.", "All devices must be on a LAN; some guest networks isolate clients; discovery can be filtered.", "Primary"),
        ("Phone hotspot", "Excellent fallback", "No external router or internet required; standard LAN sockets after guests join the hotspot.", "Users must switch Wi-Fi; behavior varies by device/carrier; host/client isolation must be tested.", "Ship with LAN"),
        ("Local-only hotspot API", "Good later", "App can create a no-internet access point with guided UX.", "Requires native Android bridge, runtime nearby-Wi-Fi permission, lifecycle handling, and OEM testing.", "Phase 2"),
        ("Wi-Fi Direct", "Good but complex", "Direct nearby connection without router; higher range and bandwidth than Bluetooth.", "Native Android API, several permissions, location-mode edge cases, device support variation; not exposed by LÖVE Lua.", "Phase 3 / optional"),
        ("Bluetooth Classic", "Weak", "Works without Wi-Fi infrastructure and can be familiar to users.", "Native bridge, pairing UX, Android permission matrix, lower throughput, cross-platform limitations, multi-peer complexity.", "Do not prioritize"),
        ("Bluetooth Low Energy", "Poor", "Low power and nearby-device discovery.", "Designed for small messages, not continuous shared-world replication; fragmentation and GATT complexity.", "Reject for gameplay"),
        ("Wi-Fi Aware / Nearby APIs", "Specialized", "Can discover/connect nearby without pre-existing LAN on supported devices.", "Native integration, support matrix, Google/API dependencies, more release risk.", "Research only if hotspot UX fails"),
    ], [1240, 1050, 2340, 3070, 1660], font_size=8.1)

    add_heading(doc, "4. Recommended player experience")
    player_flow_num = None
    add_number(doc, "Choose Local Play on the title screen, then Host Shop or Join Shop.", player_flow_num)
    add_number(doc, "The host selects a local save, display name, and 2-4 player limit, then opens a lobby.", player_flow_num)
    add_number(doc, "Joining phones list LAN shops with host name, player count, and compatible protocol status.", player_flow_num)
    add_number(doc, "A guest taps a lobby. If discovery is blocked, the host provides an IP/port fallback; a short-lived token prevents accidental joins.", player_flow_num)
    add_number(doc, "The host starts. Each player appears as a separate worker with an independent camera and controls.", player_flow_num)
    add_number(doc, "A disconnected guest gets a brief reconnect window. If the host leaves, the first release saves and closes cleanly instead of migrating the host.", player_flow_num)
    add_body(doc, "A normal phone hotspot satisfies the 'same area, no internet required' goal without adding a custom Wi-Fi stack. The lobby should explain this plainly: if no shared Wi-Fi is available, one player turns on a hotspot and the others join it before opening Local Play.")

    add_heading(doc, "4.1 Cellphone hotspot as the LAN source", level=2, page_break_before=True)
    add_callout(doc, "Preferred topology", "The phone that creates the hotspot should also host the game and own the save. Other phones join that hotspot as clients. This gives one clear session owner and avoids transferring authoritative save data between devices.", PALE_GREEN)
    add_body(doc, "There are two ways to provide this experience. The fastest MVP uses the phone's normal system Hotspot/Tethering screen: the host turns it on, guests join the displayed SSID, and the game uses ordinary local sockets. The game does not require internet access even if the hotspot also shares cellular data. This path needs only Android's INTERNET permission for the game sockets because the user, not the app, creates the hotspot.")
    add_body(doc, "A more integrated second phase can use Android's LocalOnlyHotspot API (available from API level 26). Android explicitly describes this network as a way for co-located connected devices to communicate, with no internet access. The app receives a hotspot reservation and configuration, displays the SSID/passphrase, and must close the reservation when hosting ends. This requires a small native Android-to-Lua bridge and runtime Wi-Fi permission handling.")
    add_table(doc, ["Hotspot mode", "Setup", "Permissions / code", "Recommendation"], [
        ("System hotspot / tethering", "Host opens Android settings and enables hotspot; guests join it, then return to the game.", "Game restores INTERNET only. No hotspot-control API or nearby-device prompt.", "Use for the first playable release."),
        ("App-created LocalOnlyHotspot", "Host taps Create Game Hotspot; app waits for Android callback, then shows SSID and password.", "Native wrapper bridge; ACCESS_WIFI_STATE; NEARBY_WIFI_DEVICES on Android 13+, with older-version location compatibility.", "Add after core LAN play is stable."),
    ], [1800, 2940, 3000, 1620], font_size=8.4)
    add_body(doc, "Hotspot-specific network handling:", bold_lead="Hotspot-specific network handling:")
    add_bullet(doc, "Bind the ENet host to all local interfaces, but advertise/connect through the hotspot LAN address rather than the mobile-data interface.")
    add_bullet(doc, "Prefer automatic discovery. For manual fallback, display the resolved hotspot-side host address and port; never assume a fixed address such as 192.168.43.1 because OEMs may choose another subnet.")
    add_bullet(doc, "Keep hosting in the foreground. If Android pauses, revokes, or stops the hotspot, save the shop, notify clients, and close the session cleanly.")
    add_bullet(doc, "Treat client-to-client isolation and blocked broadcast as expected network variations. Gameplay clients only need to reach the host; manual join must remain available if discovery is filtered.")
    add_bullet(doc, "Do not require cellular service. Include an explicit offline test with SIM/mobile data disabled to prove that all gameplay traffic stays local.")
    add_body(doc, "Recommended hotspot join flow:", bold_lead="Recommended hotspot join flow:")
    hotspot_flow_num = None
    add_number(doc, "Host selects Local Play > Host Shop > Use This Phone's Hotspot.", hotspot_flow_num)
    add_number(doc, "For the MVP, the game opens or explains Android hotspot settings. For the integrated version, it requests LocalOnlyHotspot and waits for onStarted.", hotspot_flow_num)
    add_number(doc, "Host keeps The Picture Shop open and shows the network name/password plus a clear Waiting for Players screen.", hotspot_flow_num)
    add_number(doc, "Guests join the hotspot in Android Wi-Fi settings, return to The Picture Shop, and select the discovered lobby or manual address.", hotspot_flow_num)
    add_number(doc, "The host verifies the player list and starts. When the session ends, the host save is committed and an app-created hotspot reservation is released.", hotspot_flow_num)

    add_heading(doc, "5. Recommended technical architecture")
    add_callout(doc, "Authority flow", "Guest input -> validated command -> host simulation -> authoritative event/snapshot -> all clients -> interpolation and rendering", PALE_BLUE)
    add_body(doc, "Use a host-authoritative model. The host runs the complete shop simulation and is the only device allowed to make durable changes. Guests do not send final values such as 'money = 5000' or 'job complete'; they send intentions such as move, interact, start machine, submit quote, or unload pallet. The host validates the request against current state, applies it once, and broadcasts the result.")

    add_heading(doc, "5.1 Transport and channels", level=2)
    add_body(doc, "Use the ENet module included with LÖVE 11.5. ENet runs over UDP but supplies optional reliable, ordered delivery, connection state, peer management, and multiple channels. That avoids building a custom reliability layer while keeping latency appropriate for movement.")
    add_table(doc, ["Channel", "Delivery", "Examples", "Suggested rate"], [
        ("0 - session", "Reliable, ordered", "hello, authentication token, ready state, join/leave, full snapshot", "Event driven"),
        ("1 - commands", "Reliable, ordered", "interact, operate machine, quote, purchase, unload, screen action", "Event driven"),
        ("2 - movement", "Unreliable / sequenced", "input vector, facing, client input sequence", "20 Hz input"),
        ("3 - snapshots", "Unreliable / sequenced", "player transforms, short-lived motion state", "10-15 Hz"),
        ("4 - durable events", "Reliable, ordered", "money/inventory delta, job transition, machine lock, message", "Event driven"),
    ], [1450, 1700, 3600, 2610], font_size=8.2)
    add_body(doc, "Start with conservative limits: maximum four peers, bounded message sizes, a fixed protocol version, and a 30-60 Hz local render loop. Smooth remote avatars with interpolation; do not require deterministic lockstep because the current simulation uses timers, mutable module state, and real-time visitor/machine updates.")

    add_heading(doc, "5.2 Discovery and connection", level=2)
    add_body(doc, "For the first release, use a small non-blocking LuaSocket UDP discovery service on a separate port. A host advertises a compact record containing only protocol version, lobby identifier, host display name, player count, ENet port, and a short-lived nonce. Clients listen for advertisements and then connect through ENet. Discovery must stop when the lobby closes or the app loses focus.")
    add_body(doc, "Keep two fallbacks because multicast/broadcast may be blocked by access-point settings: refresh the lobby list, and manual host address entry. Android Network Service Discovery (DNS-SD/mDNS) is a good native enhancement if field testing shows broadcast discovery is unreliable, but it is not required to prove the gameplay model.")

    add_heading(doc, "5.3 State ownership and conflict rules", level=2)
    add_table(doc, ["State / action", "Owner", "Conflict policy"], [
        ("Business clock, bills, emails, visitors, trucks", "Host", "Guests receive events/snapshots; never advance these systems independently."),
        ("Money, inventory, jobs, machine condition", "Host", "Atomic command validation; reject stale or impossible requests with a visible reason."),
        ("Player movement", "Host authoritative with client prediction", "Host corrects invalid movement; client reconciles to acknowledged input sequence."),
        ("Machine or pallet-jack operation", "Host lease/lock", "First valid operator owns the device until release, completion, timeout, or disconnect."),
        ("Panel/menu state", "Local client", "The visible panel is local; actions that change the shop become host commands."),
        ("Save slot", "Host only", "Autosave after durable transactions and at session close; guests never overwrite the shared shop."),
    ], [2600, 2500, 4260], font_size=8.7)

    add_heading(doc, "5.4 Safe protocol design", level=2)
    add_bullet(doc, "Use a dedicated JSON or compact binary codec; never execute received text as Lua.")
    add_bullet(doc, "Every packet carries protocolVersion, messageType, sessionId, senderId, and sequence or event ID where applicable.")
    add_bullet(doc, "Validate allowed keys, numeric ranges, string length, array count, packet size, action rate, and the sender's current permissions/lease.")
    add_bullet(doc, "Use stable identifiers for players, jobs, pallets, machines, shipments, and commands. Make durable commands idempotent so retries cannot purchase or pay twice.")
    add_bullet(doc, "Send a full, validated join snapshot, then deltas/events. Periodically send a compact checksum or state revision to detect desynchronization.")
    add_bullet(doc, "Reject version/build mismatches with a clear UI message before loading the shared shop.")

    add_heading(doc, "6. Codebase changes")
    add_table(doc, ["Area", "Proposed change", "Why"], [
        ("src/net/protocol.lua", "Message constants, schema validators, size/range limits, stable serialization.", "Creates a safe boundary separate from local saves."),
        ("net/transport_enet.lua", "Host/client creation, channel policy, polling, disconnects, statistics.", "Shared cross-platform session transport."),
        ("src/net/discovery.lua", "Non-blocking UDP advertise/listen and manual-address fallback.", "Find nearby lobbies on ordinary LAN/hotspot."),
        ("src/net/session.lua", "Lobby, player IDs, handshake, reconnect token, protocol/build checks.", "Owns multiplayer lifecycle without crowding app.lua."),
        ("src/net/replication.lua", "Full snapshot, delta/event application, revisions, interpolation buffers.", "Keeps clients synchronized."),
        ("src/commands/*.lua", "Validated command handlers around existing domain functions.", "Moves mutation authority to the host."),
        ("src/state.lua / save schema", "Add participants and multiplayer-safe snapshot fields; preserve migration from version 12.", "Represent several players while keeping old saves loadable."),
        ("src/world.lua", "Replace World.player singleton with participant-aware lookup/rendering; host collision authority.", "Supports several independent avatars."),
        ("src/app.lua", "Create/update session service, route input as local or network command, save only on host.", "Integrates networking at the application boundary."),
        ("Title/HUD screens", "Local Play lobby, connection status, player labels, reconnect/host-ended messages.", "Makes the feature understandable on phones."),
        ("Android manifest/build", "Restore android.permission.INTERNET and verify merged permissions in the APK report.", "Android requires this permission for network sockets, including LAN sockets."),
    ], [1850, 4620, 2890], font_size=8.2)

    add_callout(doc, "Architecture principle", "Keep transport, discovery, and Android permission handling at the edges. Keep money, jobs, machines, pallets, and business rules in shared domain modules so desktop and Android continue to ship from one source tree.", PALE_GREEN)

    add_heading(doc, "7. Android permissions and platform work")
    add_body(doc, "Ordinary same-LAN/hotspot sockets need android.permission.INTERNET. Despite the name, this is the Android permission used for standard network sockets; the current manifest explicitly removes it and must be changed. It is a normal install-time permission and does not itself mean the game needs a cloud service.")
    add_body(doc, "Do not add Bluetooth, location, or nearby-Wi-Fi permissions for the basic shared-Wi-Fi implementation. Add permissions only when a selected feature actually needs them:")
    add_bullet(doc, "Wi-Fi Direct or local-only hotspot on Android 13+: NEARBY_WIFI_DEVICES; relevant Wi-Fi state/change permissions; older Android paths may require location permission depending on API used.")
    add_bullet(doc, "Bluetooth discovery/connection on Android 12+: BLUETOOTH_SCAN and/or BLUETOOTH_CONNECT, plus runtime Nearby devices approval.")
    add_bullet(doc, "Future Android local-network protection: Android's current guidance says apps targeting Android 17 / SDK 37 or later need ACCESS_LOCAL_NETWORK for broad direct LAN access, while older target SDKs temporarily receive access through INTERNET. Re-check this at the target-SDK upgrade because the platform policy is evolving.")
    add_body(doc, "If a native discovery bridge is later added, keep it in the existing Android wrapper boundary and expose only small lifecycle-safe functions to Lua: start advertising, stop advertising, start discovery, stop discovery, and return resolved host/port records.")

    add_heading(doc, "8. Delivery roadmap and effort")
    add_body(doc, "The ranges below assume one experienced engineer already familiar with this codebase, with Android devices available for testing. Art production, new character animation, store review time, and internet multiplayer are excluded.")
    add_table(doc, ["Phase", "Deliverable", "Estimated effort", "Exit criterion"], [
        ("0 - spike", "Two Android builds connect over ENet; permission, bind, send/receive, hotspot, suspend/resume tests.", "3-5 days", "Round-trip messages work on two physical phones."),
        ("1 - architecture", "Command boundary, participant model, safe protocol, host-only save path, loopback harness.", "2-4 weeks", "Two desktop instances stay synchronized through core actions."),
        ("2 - playable MVP", "Lobby/discovery, 2-4 avatars, movement, interaction locks, core job/machine/pallet flows, reconnect handling.", "4-6 weeks", "Two Android phones can complete a representative shop session."),
        ("3 - hardening", "Packet-loss/reorder tests, compatibility errors, UI polish, full gameplay audit, device/router matrix, recovery.", "3-5 weeks", "Release checklist passes across representative networks/devices."),
        ("Optional - native discovery", "Android NSD or local-only hotspot bridge.", "+2-4 weeks", "Reliable no-manual-address connection on supported devices."),
        ("Optional - Wi-Fi Direct", "Peer/service discovery, group creation, permissions, socket handoff, OEM matrix.", "+4-8 weeks", "No-router direct connection is reliable on supported devices."),
        ("Optional - Bluetooth", "Native transport, permissions, pairing, fragmentation, multi-peer and lifecycle handling.", "+6-10 weeks", "Only pursue if a product requirement outweighs its cost."),
    ], [1200, 3900, 1300, 2960], font_size=8.2)
    add_callout(doc, "Planning range", "A credible Android LAN release is approximately 9-15 engineer-weeks after a successful spike. The uncertainty is mainly the single-player-to-authoritative-command refactor and full gameplay synchronization, not raw socket connectivity.", PALE_BLUE)

    add_heading(doc, "9. Verification plan")
    add_heading(doc, "9.1 Automated tests", level=2)
    add_bullet(doc, "Protocol round-trip and rejection tests: unknown type, oversized packet, invalid range, invalid UTF-8/text length, duplicate command, stale revision, wrong session/version.")
    add_bullet(doc, "Two-instance deterministic scenarios: join, move, use machine, purchase, accept/decline/complete job, unload truck, relocate machine, operate pallet jack, save, quit, reload.")
    add_bullet(doc, "Network simulation: latency, jitter, 1-10% loss, duplication, reordering, reconnect, host termination, client pause/resume, and Wi-Fi handoff.")
    add_bullet(doc, "State invariants after every durable event: nonnegative inventory, single pallet owner/location, no double payment, one active machine operator, monotonically increasing revision.")
    add_bullet(doc, "Backward compatibility: existing version-12 single-player saves migrate and remain playable offline.")
    add_heading(doc, "9.2 Physical test matrix", level=2)
    add_table(doc, ["Network", "Devices", "Required checks"], [
        ("Home router, 2.4 GHz", "Two Android phones", "Discover, join, 30-minute play, suspend/resume, reconnect."),
        ("Home router, 5/6 GHz", "Three or four phones", "Mixed device ages, lobby capacity, bandwidth and heat."),
        ("System phone hotspot", "Hotspot owner also hosts + 1-3 guests", "Mobile data on/off, discovery, correct hotspot IP, manual fallback, suspend/stop behavior."),
        ("System phone hotspot", "Hotspot owner only; second phone hosts", "Guest-to-guest reachability and AP client-isolation behavior; document if unsupported."),
        ("LocalOnlyHotspot API", "Host + 1-3 guests on Android 8+", "Permission grant/deny, SSID/passphrase, onStarted/onStopped, reservation cleanup, no internet."),
        ("Guest/corporate Wi-Fi", "Two phones", "Client isolation produces a clear failure and fallback guidance."),
        ("Windows + Android", "Desktop host/client plus phone", "Shared protocol and build compatibility, if cross-device play is in scope."),
    ], [2000, 2050, 5310], font_size=8.7)
    add_body(doc, "Track median and 95th-percentile round-trip time, correction distance, packet loss, bytes per second, disconnect reason, join time, and state revision gaps. Keep diagnostic logs bounded and free of save contents or personal network identifiers.")

    add_heading(doc, "10. Major risks and mitigations")
    add_table(doc, ["Risk", "Impact", "Mitigation"], [
        ("Independent simulation diverges", "High", "Host-only simulation; clients render replicated state and never advance durable systems independently."),
        ("Direct state mutations bypass commands", "High", "Incrementally wrap mutation entry points; add audit tests that fail when guests save or mutate authoritative state."),
        ("Discovery blocked by router/OEM", "Medium", "Manual address fallback; test NSD/native enhancement only if data justifies it."),
        ("Two players use one machine/pallet", "High", "Host-issued leases with timeout, release, and disconnect cleanup."),
        ("Duplicate economic action", "High", "Command IDs, idempotent handlers, host-side revision checks, atomic apply/broadcast/save."),
        ("Host app backgrounds or closes", "Medium", "Autosave durable actions, pause networking cleanly, notify clients, bounded reconnect window."),
        ("Permission/policy change", "Medium", "Verify merged manifest and target SDK every release; test Android local-network restrictions before target upgrades."),
        ("Large snapshots / mobile heat", "Medium", "Bounded 10-15 Hz transforms, event deltas, compression only after measurement, profiling on physical devices."),
    ], [2600, 1250, 5510], font_size=8.4)

    add_heading(doc, "11. Go / no-go recommendation")
    add_body(doc, "Proceed with a short ENet-on-Android spike, then build the shared-Wi-Fi/hotspot version if it passes. This is a strong fit for the engine and the same-source mobile architecture, and it directly meets the requirement that nearby players can play together without an external server.")
    add_body(doc, "Do not begin with Bluetooth or Wi-Fi Direct. Both would add native Android work before the game has a multiplayer authority model, and neither solves the harder synchronization problem. Establish the host-authoritative command and replication layer first; alternate radios can then reuse it later.")
    add_body(doc, "Recommended acceptance target for the first release: two to four Android players can discover or manually join a host on shared Wi-Fi or a phone hotspot, play a 30-minute shared-shop session, operate representative machines without duplication, reconnect one guest, and leave with a valid host save that loads correctly offline.")

    add_heading(doc, "Appendix A - Proposed first spike")
    spike_num = None
    add_number(doc, "Restore android.permission.INTERNET in the embedded Android manifest and add a build assertion that the merged APK contains it.", spike_num)
    add_number(doc, "Create a minimal host with enet.host_create('*:port') and a minimal client that connects to the displayed host address; poll with zero timeout from love.update so the render loop never blocks.", spike_num)
    add_number(doc, "Exchange protocol/build hello messages, a ping/pong with timestamps, and a small reliable state record. Log round-trip time and disconnect reason.", spike_num)
    add_number(doc, "Test two physical phones on the same router and with the hotspot owner also hosting the game. Repeat with mobile data and SIM disabled, then test focus loss, app switching, reconnect, and hotspot stop.", spike_num)
    add_number(doc, "Prove UDP discovery separately with LuaSocket. If broadcast fails on a network, confirm that manual address entry still connects.", spike_num)
    add_number(doc, "Record results in an engineering decision: transport viability, permission behavior, device/network matrix, observed latency/loss, and whether native discovery is justified.", spike_num)

    add_heading(doc, "Appendix B - Sources and evidence")
    add_body(doc, "Repository evidence reviewed:", bold_lead="Repository evidence reviewed:")
    for item in (
        "MOBILE_ARCHITECTURE_DECISION.md; ANDROID_PORT.md; conf.lua; main.lua",
        "src/app.lua; src/state.lua; src/world.lua; src/input.lua; src/save.lua; src/save_schema.lua",
        "mobile/android/AndroidManifest.xml; mobile/config.json; tools/build_android_apk.ps1",
        "Staged LÖVE 11.5 Android source under output/mobile/love-android, including ENet and LuaSocket preload code",
    ):
        add_bullet(doc, item)

    add_body(doc, "Official platform references checked August 25, 2026:", bold_lead="Official platform references checked August 25, 2026:")
    references = [
        ("LÖVE wiki - lua-enet (included with LÖVE; reliable networking over UDP)", "https://love2d.org/wiki/lua-enet"),
        ("LÖVE wiki - socket (LuaSocket TCP/UDP; non-blocking guidance)", "https://love2d.org/wiki/socket"),
        ("Android Developers - Connect devices wirelessly", "https://developer.android.com/develop/connectivity/wifi"),
        ("Android Developers - Use network service discovery", "https://developer.android.com/develop/connectivity/wifi/use-nsd"),
        ("Android Developers - Create P2P connections with Wi-Fi Direct", "https://developer.android.com/develop/connectivity/wifi/wifi-direct"),
        ("Android Developers - Request permission to access nearby Wi-Fi devices", "https://developer.android.com/develop/connectivity/wifi/wifi-permissions"),
        ("Android Developers - Use a local-only Wi-Fi hotspot", "https://developer.android.com/develop/connectivity/wifi/localonlyhotspot"),
        ("Android Developers - Local network permission", "https://developer.android.com/privacy-and-security/local-network-permission"),
        ("Android Developers - Bluetooth permissions", "https://developer.android.com/develop/connectivity/bluetooth/bt-permissions"),
    ]
    for title_text, url in references:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(5)
        add_hyperlink(p, title_text, url)

    doc.core_properties.title = "Local Multiplayer Feasibility Report - The Picture Shop"
    doc.core_properties.subject = "Android LAN, hotspot, Wi-Fi Direct, and Bluetooth feasibility"
    doc.core_properties.author = "The Picture Shop development team"
    doc.core_properties.keywords = "local multiplayer, LAN, Android, LÖVE, ENet, Wi-Fi, Bluetooth"
    doc.core_properties.comments = "Prepared from repository inspection and official platform documentation."
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
